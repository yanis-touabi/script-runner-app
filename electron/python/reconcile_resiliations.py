"""Move cancelled SIM cards from ACTIVES TOTAL to resiliation sheets."""

from __future__ import annotations

import math
import re
import time
import unicodedata
from bisect import bisect_left
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path
from typing import Any

import openpyxl


OPERATORS = ("MOBILIS", "DJEZZY", "OOREDOO")
SOURCE_SHEET = "ACTIVES TOTAL"
DESTINATION_SHEETS = ("SIM EN COURS DE Résiliation", "Résiliation En Cours")
RESILIATION_SERIAL_COLUMN = "Numero SIM Abrege"
ACTIVE_SERIAL_COLUMN = "Num Serie"
OUTPUT_COLUMNS = ("Numero SIM Abrege", "Is Found", "Active File", "Active Row")


def _text(value: Any) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return text


def normalize_serial(value: Any) -> str:
    """Return the 12-digit SIM core used for all comparisons."""
    text = _text(value)
    if not text.isdigit():
        return ""
    if len(text) == 12:
        return text
    if len(text) == 18:
        return text[6:]
    if len(text) == 19:
        return text[6:-1]
    return ""


def _header_key(value: Any) -> str:
    text = unicodedata.normalize("NFKD", _text(value))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"\s+", " ", text).strip().casefold()


def _find_folder(root: Path, name: str) -> Path:
    matches = [path for path in root.rglob("*") if path.is_dir() and path.name.casefold() == name.casefold()]
    if len(matches) != 1:
        raise FileNotFoundError(f"Expected one {name!r} folder below {root}, found {len(matches)}.")
    return matches[0]


def _operator_for_file(path: Path) -> str | None:
    stem = path.stem.casefold()
    return next((operator for operator in OPERATORS if operator.casefold() in stem), None)


def _workbooks(folder: Path) -> list[Path]:
    return sorted(
        path
        for path in folder.rglob("*")
        if path.is_file()
        and path.suffix.casefold() in {".xlsx", ".xlsm"}
        and not path.name.startswith("~$")
    )


def _header_map(sheet: Any) -> dict[str, int]:
    headers = {}
    for cell in sheet[1]:
        header = _header_key(cell.value)
        if header:
            headers[header] = cell.column
    return headers


def _sheet_by_name(workbook: Any, expected_names: tuple[str, ...]) -> Any | None:
    expected = {_header_key(name) for name in expected_names}
    return next((sheet for sheet in workbook.worksheets if _header_key(sheet.title) in expected), None)


def _resiliation_keys(resiliation_root: Path) -> dict[str, set[str]]:
    keys: dict[str, set[str]] = defaultdict(set)
    for path in _workbooks(resiliation_root):
        operator = _operator_for_file(path)
        if operator is None:
            continue
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                headers = _header_map(sheet)
                serial_column = headers.get(_header_key(RESILIATION_SERIAL_COLUMN))
                if serial_column is None:
                    raise ValueError(f"Missing {RESILIATION_SERIAL_COLUMN!r} in {path} [{sheet.title}]")
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    value = row[serial_column - 1] if serial_column <= len(row) else None
                    normalized = normalize_serial(value)
                    if normalized:
                        keys[operator].add(normalized)
        finally:
            workbook.close()
    return keys


def _copy_row_style(
    sheet: Any, source_row: int, destination_row: int, columns: list[int]
) -> None:
    for column in columns:
        source_cell = sheet.cell(source_row, column)
        destination_cell = sheet.cell(destination_row, column)
        if source_cell.has_style:
            destination_cell._style = copy(source_cell._style)
        destination_cell.number_format = source_cell.number_format
    if sheet.row_dimensions[source_row].height is not None:
        sheet.row_dimensions[destination_row].height = sheet.row_dimensions[source_row].height


def _last_data_row(sheet: Any, headers: dict[str, int]) -> int:
    for row_number in range(sheet.max_row, 1, -1):
        if any(sheet.cell(row_number, column).value is not None for column in headers.values()):
            return row_number
    return 1


def _compact_rows(sheet: Any, rows_to_remove: list[int], max_column: int) -> None:
    """Remove rows in one pass, avoiding repeated worksheet-wide row shifts."""
    removed = sorted(rows_to_remove)
    removed_set = set(removed)
    new_cells = {}
    for (row_number, column), cell in sheet._cells.items():
        if column > max_column or row_number in removed_set:
            continue
        new_row = row_number - bisect_left(removed, row_number)
        cell.row = new_row
        new_cells[(new_row, column)] = cell
    sheet._cells = new_cells

    for row_number in list(sheet.row_dimensions):
        if row_number in removed_set:
            del sheet.row_dimensions[row_number]
        else:
            new_row = row_number - bisect_left(removed, row_number)
            if new_row != row_number:
                dimension = sheet.row_dimensions.pop(row_number)
                dimension.index = new_row
                sheet.row_dimensions[new_row] = dimension


def _append_record(
    source_sheet: Any,
    source_row: int,
    destination_sheet: Any,
    source_headers: dict[str, int],
    destination_headers: dict[str, int],
    destination_row: int,
    style_row: int,
) -> int:
    for header, destination_column in destination_headers.items():
        source_column = source_headers.get(header)
        value = source_sheet.cell(source_row, source_column).value if source_column else None
        destination_sheet.cell(destination_row, destination_column).value = value
    if style_row > 1:
        _copy_row_style(destination_sheet, style_row, destination_row, list(destination_headers.values()))
    return destination_row + 1


def _move_workbook(path: Path, operator: str, keys: set[str], root: Path) -> tuple[list[dict[str, str]], int]:
    started = time.perf_counter()
    print(f"[{operator}] Reading {path.relative_to(root)}...", flush=True)
    workbook = openpyxl.load_workbook(path, keep_vba=path.suffix.casefold() == ".xlsm")
    source_sheet = _sheet_by_name(workbook, (SOURCE_SHEET,))
    if source_sheet is None:
        workbook.close()
        return [], 0
    destination_sheet = _sheet_by_name(workbook, DESTINATION_SHEETS)
    if destination_sheet is None:
        workbook.close()
        raise ValueError(f"Missing destination resiliation sheet in {path}")

    source_headers = _header_map(source_sheet)
    destination_headers = _header_map(destination_sheet)
    next_destination_row = _last_data_row(destination_sheet, destination_headers) + 1
    style_row = next_destination_row - 1
    serial_column = source_headers.get(_header_key(ACTIVE_SERIAL_COLUMN))
    if serial_column is None:
        workbook.close()
        raise ValueError(f"Missing {ACTIVE_SERIAL_COLUMN!r} in {path} [{source_sheet.title}]")

    matches: list[tuple[int, str]] = []
    for row_number in range(2, source_sheet.max_row + 1):
        normalized = normalize_serial(source_sheet.cell(row_number, serial_column).value)
        if normalized in keys:
            matches.append((row_number, normalized))
    print(f"[{operator}] Found {len(matches)} matching row(s) in {path.relative_to(root)}", flush=True)

    locations: list[dict[str, str]] = []
    for row_number, normalized in matches:
        destination_row = next_destination_row
        next_destination_row = _append_record(
            source_sheet,
            row_number,
            destination_sheet,
            source_headers,
            destination_headers,
            next_destination_row,
            style_row,
        )
        locations.append(
            {
                "serial": normalized,
                "file": str(path.relative_to(root)),
                "row": str(destination_row),
            }
        )

    if matches:
        _compact_rows(
            source_sheet,
            [row_number for row_number, _ in matches],
            max(source_headers.values()),
        )

    if matches:
        print(f"[{operator}] Saving {len(matches)} moved row(s) to {path.relative_to(root)}...", flush=True)
        workbook.save(path)
    workbook.close()
    elapsed = time.perf_counter() - started
    print(f"[{operator}] Finished {path.relative_to(root)} in {elapsed:.1f}s", flush=True)
    return locations, len(matches)


def _write_found_files(root: Path, keys: dict[str, set[str]], locations: dict[str, list[dict[str, str]]]) -> None:
    for operator in OPERATORS:
        output = openpyxl.Workbook()
        sheet = output.active
        sheet.title = "Résiliations"
        sheet.append(list(OUTPUT_COLUMNS))
        for serial in sorted(keys.get(operator, set())):
            matches = [location for location in locations[operator] if location["serial"] == serial]
            sheet.append(
                [
                    serial,
                    "Found" if matches else "Not Found",
                    "; ".join(location["file"] for location in matches),
                    "; ".join(location["row"] for location in matches),
                ]
            )
        output.save(root / f"{operator}_resiliations_found.xlsx")
        output.close()


def main() -> None:
    started = time.perf_counter()
    root = Path(__file__).resolve().parent
    sim_db_root = _find_folder(root, "SIM_DB")
    resiliation_root = _find_folder(root, "RESILIATION")
    print("Reading resiliation files...", flush=True)
    keys = _resiliation_keys(resiliation_root)
    locations: dict[str, list[dict[str, str]]] = defaultdict(list)
    moved = Counter()
    skipped: list[str] = []

    for path in _workbooks(sim_db_root):
        operator = _operator_for_file(path)
        if operator is None:
            continue
        file_locations, count = _move_workbook(path, operator, keys[operator], root)
        locations[operator].extend(file_locations)
        moved[operator] += count
        if count == 0:
            skipped.append(str(path.relative_to(root)))

    _write_found_files(root, keys, locations)
    for operator in OPERATORS:
        print(
            f"{operator}: moved {moved[operator]} record(s), "
            f"active matches {len(locations[operator])}, "
            f"resiliation keys {len(keys[operator])}"
        )
    if skipped:
        print("No matching rows in: " + ", ".join(skipped))
    print(f"Completed in {time.perf_counter() - started:.1f}s", flush=True)


if __name__ == "__main__":
    main()
